#!/usr/bin/env python3
"""
Проверка списка УИН ювелирных изделий в ГИИС ДМДК (Федеральная пробирная палата).

Скрипт открывает публичный сервис проверки УИН (по умолчанию https://dmdk.ru/check-uin/)
в браузере Chromium, по очереди вводит каждый УИН, собирает ответ сервиса и сохраняет
сводку в Excel (.xlsx) и CSV: статус изделия (в обороте / продано / выведено из оборота /
не найдено), а также все поля, которые показал сервис (проба, металл, изготовитель и т.д.).

Как именно устроена страница сервиса, скрипт заранее не знает — он:
  1) перехватывает JSON-ответы, которые страница получает от сервера после ввода УИН;
  2) дополнительно берёт текст, появившийся на странице после проверки.
Поэтому небольшие изменения вёрстки сайта его не ломают. Если сайт изменится сильнее —
селекторы поля ввода и кнопки можно задать параметрами --input-selector / --button-selector.

Запуск:
    python dmdk_checker.py uins.txt                 # результат в results_<дата>.xlsx
    python dmdk_checker.py uins.xlsx -o итог.xlsx --show-browser
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import time
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Callable, Iterable

DEFAULT_URL = "https://dmdk.ru/check-uin/"

# Статусы, которые мы выводим в итоговую таблицу. Порядок правил важен:
# «выведено из оборота» проверяется раньше «в обороте», «не найден» — раньше всего.
STATUS_NOT_FOUND = "Не найден"
STATUS_SOLD = "Продано"
STATUS_WITHDRAWN = "Выведено из оборота"
STATUS_IN_CIRCULATION = "В обороте"
STATUS_UNKNOWN = "Не распознан"
STATUS_ERROR = "Ошибка проверки"

STATUS_RULES: list[tuple[str, tuple[str, ...]]] = [
    (STATUS_NOT_FOUND, (
        "не найден", "не найдено", "не зарегистрирован", "отсутству", "нет данных",
        "not found", "не существует",
    )),
    (STATUS_SOLD, (
        "продан", "продажа", "реализован", "розничная продажа", "sold", "продано",
    )),
    (STATUS_WITHDRAWN, (
        "выведен из оборота", "выведено из оборота", "вывод из оборота", "выбыл",
        "списан", "утилизирован", "переработ", "аффинаж", "withdrawn",
    )),
    (STATUS_IN_CIRCULATION, (
        "в обороте", "введен в оборот", "введено в оборот", "в наличии", "на складе",
        "in circulation", "зарегистрирован",
    )),
]

# Ключи JSON, в которых обычно лежит статус изделия.
STATUS_KEYS = ("status", "state", "statusname", "statusdescription", "статус", "состояние")

INPUT_SELECTORS = [
    "input[name*='uin' i]",
    "input[id*='uin' i]",
    "input[placeholder*='УИН' i]",
    "input[placeholder*='uin' i]",
    "input[type='search']",
    "input[type='text']",
    "input:not([type])",
]
BUTTON_TEXTS = ["Проверить", "Найти", "Поиск", "Искать", "Проверка", "Check", "Search"]


@dataclass
class CheckResult:
    uin: str
    status: str = STATUS_UNKNOWN
    status_raw: str = ""
    details: dict[str, str] = field(default_factory=dict)
    page_text: str = ""
    raw_json: str = ""
    error: str = ""
    checked_at: str = field(default_factory=lambda: datetime.now().strftime("%Y-%m-%d %H:%M:%S"))


# ---------------------------------------------------------------------------
# Чтение списка УИН
# ---------------------------------------------------------------------------

def normalize_uin(value: object) -> str:
    """Убирает пробелы/дефисы; числа из Excel приводит к строке без «.0» и экспоненты."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    text = re.sub(r"[\s\-]", "", text)
    return text


def looks_like_uin(text: str) -> bool:
    # УИН ювелирного изделия — цифровой код (обычно 16 цифр). Берём с запасом.
    return bool(re.fullmatch(r"\d{8,40}", text))


def read_uins(path: Path) -> list[str]:
    """Читает УИН из .txt / .csv / .xlsx. Берёт все ячейки, похожие на УИН, без дублей."""
    raw: list[object] = []
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                raw.extend(row)
    elif suffix == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as f:
            sample = f.read(4096)
            f.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
            except csv.Error:
                dialect = csv.excel
            for row in csv.reader(f, dialect):
                raw.extend(row)
    else:
        text = path.read_text(encoding="utf-8-sig")
        # Один УИН на строку (пробелы внутри допустимы: «6433 0000 …») или через , ; Tab.
        for line in text.splitlines():
            raw.extend(re.split(r"[,;\t]+", line))

    seen: set[str] = set()
    uins: list[str] = []
    for item in raw:
        uin = normalize_uin(item)
        if looks_like_uin(uin) and uin not in seen:
            seen.add(uin)
            uins.append(uin)
    return uins


# ---------------------------------------------------------------------------
# Разбор ответа сервиса
# ---------------------------------------------------------------------------

def flatten_json(data: object, prefix: str = "") -> dict[str, str]:
    """{"a": {"b": 1}, "c": [2, 3]} -> {"a.b": "1", "c[0]": "2", "c[1]": "3"}"""
    out: dict[str, str] = {}
    if isinstance(data, dict):
        for key, value in data.items():
            out.update(flatten_json(value, f"{prefix}.{key}" if prefix else str(key)))
    elif isinstance(data, list):
        for i, value in enumerate(data):
            out.update(flatten_json(value, f"{prefix}[{i}]"))
    elif data is not None and data != "":
        out[prefix or "value"] = str(data)
    return out


def classify_status(text: str) -> str:
    low = text.lower().replace("ё", "е")
    for status, keywords in STATUS_RULES:
        if any(k in low for k in keywords):
            return status
    return STATUS_UNKNOWN


def extract_status(details: dict[str, str], page_text: str) -> tuple[str, str]:
    """Возвращает (нормализованный статус, исходная формулировка статуса)."""
    # 1. Явное поле статуса в JSON.
    for key, value in details.items():
        last = re.sub(r"\[\d+\]", "", key.split(".")[-1]).lower()
        if last in STATUS_KEYS and value:
            status = classify_status(value)
            if status != STATUS_UNKNOWN:
                return status, value
    # 2. Строка «Статус: ...» в тексте страницы.
    m = re.search(r"(?:статус|состояние)[^\n:]*[:\n]\s*([^\n]+)", page_text, re.IGNORECASE)
    if m:
        status = classify_status(m.group(1))
        if status != STATUS_UNKNOWN:
            return status, m.group(1).strip()
    # 3. Ключевые слова где угодно (сначала в тексте страницы, потом в JSON).
    for source in (page_text, " ".join(details.values())):
        status = classify_status(source)
        if status != STATUS_UNKNOWN:
            return status, ""
    return STATUS_UNKNOWN, ""


def parse_page_fields(page_text: str) -> dict[str, str]:
    """Пары «Название: значение» из текста результата на странице."""
    fields: dict[str, str] = {}
    lines = [ln.strip() for ln in page_text.splitlines() if ln.strip()]
    for i, line in enumerate(lines):
        if ":" in line:
            name, _, value = line.partition(":")
            name, value = name.strip(), value.strip()
            if not value and i + 1 < len(lines) and ":" not in lines[i + 1]:
                value = lines[i + 1]
            if name and value and len(name) <= 60:
                fields.setdefault(name, value)
    return fields


def new_text(before: str, after: str) -> str:
    """Строки, которые появились на странице после проверки (результат проверки)."""
    old = set(ln.strip() for ln in before.splitlines())
    added = [ln.strip() for ln in after.splitlines() if ln.strip() and ln.strip() not in old]
    return "\n".join(added)


# ---------------------------------------------------------------------------
# Работа с браузером
# ---------------------------------------------------------------------------

class DmdkChecker:
    def __init__(
        self,
        url: str = DEFAULT_URL,
        headless: bool = True,
        timeout: float = 30.0,
        input_selector: str | None = None,
        button_selector: str | None = None,
        executable_path: str | None = None,
        log: Callable[[str], None] = print,
    ):
        self.url = url
        self.headless = headless
        self.timeout_ms = int(timeout * 1000)
        self.input_selector = input_selector
        self.button_selector = button_selector
        self.executable_path = executable_path
        self.log = log
        self._pw = None
        self._browser = None
        self.page = None

    def __enter__(self) -> "DmdkChecker":
        from playwright.sync_api import sync_playwright

        self._pw = sync_playwright().start()
        launch_kwargs: dict = {"headless": self.headless}
        if self.executable_path:
            launch_kwargs["executable_path"] = self.executable_path
        self._browser = self._pw.chromium.launch(**launch_kwargs)
        context = self._browser.new_context(
            locale="ru-RU",
            viewport={"width": 1280, "height": 900},
            user_agent=(
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                "(KHTML, like Gecko) Chrome/140.0 Safari/537.36"
            ),
        )
        self.page = context.new_page()
        self.page.set_default_timeout(self.timeout_ms)
        return self

    def __exit__(self, *exc) -> None:
        try:
            if self._browser:
                self._browser.close()
        finally:
            if self._pw:
                self._pw.stop()

    # -- поиск элементов формы ------------------------------------------------

    def _find_input(self):
        selectors = [self.input_selector] if self.input_selector else INPUT_SELECTORS
        for sel in selectors:
            loc = self.page.locator(sel)
            for i in range(loc.count()):
                el = loc.nth(i)
                if el.is_visible() and el.is_editable():
                    return el
        raise RuntimeError(
            "Не найдено поле ввода УИН на странице. Укажите его вручную: --input-selector"
        )

    def _submit(self, field) -> None:
        if self.button_selector:
            self.page.locator(self.button_selector).first.click()
            return
        # Кнопка внутри той же формы, затем любая видимая кнопка с подходящим текстом.
        form = field.locator("xpath=ancestor::form[1]")
        scopes = [form] if form.count() else []
        scopes.append(self.page)
        for scope in scopes:
            for text in BUTTON_TEXTS:
                btn = scope.locator(
                    f"button:has-text('{text}'), input[type=submit][value*='{text}' i], "
                    f"a:has-text('{text}'), [role=button]:has-text('{text}')"
                )
                for i in range(btn.count()):
                    if btn.nth(i).is_visible():
                        btn.nth(i).click()
                        return
            sub = scope.locator("button[type=submit], input[type=submit]")
            for i in range(sub.count()):
                if sub.nth(i).is_visible():
                    sub.nth(i).click()
                    return
        field.press("Enter")

    def _body_text(self) -> str:
        try:
            return self.page.inner_text("body", timeout=5000)
        except Exception:
            return ""

    def _wait_for_captcha(self) -> None:
        text = self._body_text().lower()
        if "captcha" in self.page.content().lower() or "не робот" in text or "капч" in text:
            if self.headless:
                self.log("  ! На странице капча. Запустите с --show-browser и введите её вручную.")
            else:
                self.log("  ! На странице капча — решите её в окне браузера, ждём до 2 минут…")
                deadline = time.time() + 120
                while time.time() < deadline:
                    time.sleep(2)
                    t = self._body_text().lower()
                    if "не робот" not in t and "капч" not in t:
                        break

    # -- проверка одного УИН --------------------------------------------------

    def check(self, uin: str) -> CheckResult:
        result = CheckResult(uin=uin)
        captured: list[object] = []

        def on_response(resp) -> None:
            try:
                if resp.request.resource_type not in ("xhr", "fetch"):
                    return
                ctype = resp.headers.get("content-type", "")
                if "json" not in ctype:
                    return
                data = resp.json()
                if uin in resp.url or uin in json.dumps(data, ensure_ascii=False) or not captured:
                    captured.append(data)
            except Exception:
                pass

        self.page.goto(self.url, wait_until="domcontentloaded")
        try:
            self.page.wait_for_load_state("networkidle", timeout=10000)
        except Exception:
            pass
        self._wait_for_captcha()

        field_el = self._find_input()
        before = self._body_text()

        self.page.on("response", on_response)
        try:
            field_el.fill("")
            field_el.type(uin, delay=20)
            self._submit(field_el)
            # Ждём, пока страница получит ответ и перерисуется.
            deadline = time.time() + self.timeout_ms / 1000
            after = before
            while time.time() < deadline:
                try:
                    self.page.wait_for_load_state("networkidle", timeout=3000)
                except Exception:
                    pass
                after = self._body_text()
                if captured or after != before:
                    time.sleep(0.7)  # даём дорисовать результат целиком
                    after = self._body_text()
                    break
                time.sleep(0.5)
        finally:
            self.page.remove_listener("response", on_response)

        result.page_text = new_text(before, after)
        details: dict[str, str] = {}
        if captured:
            data = captured[0] if len(captured) == 1 else captured
            result.raw_json = json.dumps(data, ensure_ascii=False)
            details.update(flatten_json(data))
        for name, value in parse_page_fields(result.page_text).items():
            details.setdefault(name, value)
        result.details = details
        result.status, result.status_raw = extract_status(details, result.page_text)
        if not captured and not result.page_text:
            result.error = "Сервис не вернул ответ (страница не изменилась)"
        return result


# ---------------------------------------------------------------------------
# Сохранение результатов
# ---------------------------------------------------------------------------

BASE_COLUMNS = ["УИН", "Статус", "Статус (как на сайте)", "Проверено", "Ошибка"]
TAIL_COLUMNS = ["Текст ответа", "JSON ответа"]


def result_to_row(r: CheckResult) -> dict[str, str]:
    row = {
        "УИН": r.uin,
        "Статус": r.status,
        "Статус (как на сайте)": r.status_raw,
        "Проверено": r.checked_at,
        "Ошибка": r.error,
    }
    for key, value in r.details.items():
        # Поле «Статус» со страницы не должно затирать наш нормализованный статус.
        if key in BASE_COLUMNS or key in TAIL_COLUMNS:
            key = f"{key} (сайт)"
        row.setdefault(key, value)
    row["Текст ответа"] = r.page_text
    row["JSON ответа"] = r.raw_json
    return row


def collect_columns(rows: Iterable[dict[str, str]]) -> list[str]:
    cols = list(BASE_COLUMNS)
    for row in rows:
        for key in row:
            if key not in cols and key not in TAIL_COLUMNS:
                cols.append(key)
    return cols + TAIL_COLUMNS


def save_results(results: list[CheckResult], out_path: Path) -> list[Path]:
    rows = [result_to_row(r) for r in results]
    cols = collect_columns(rows)
    written: list[Path] = []

    csv_path = out_path.with_suffix(".csv")
    with csv_path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=cols, delimiter=";", extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    written.append(csv_path)

    if out_path.suffix.lower() == ".xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        colors = {
            STATUS_IN_CIRCULATION: "C6EFCE",
            STATUS_SOLD: "BDD7EE",
            STATUS_WITHDRAWN: "FFEB9C",
            STATUS_NOT_FOUND: "FFC7CE",
            STATUS_ERROR: "FFC7CE",
            STATUS_UNKNOWN: "EDEDED",
        }
        wb = Workbook()
        ws = wb.active
        ws.title = "Результаты"
        ws.append(cols)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append([row.get(c, "") for c in cols])
            fill = colors.get(row["Статус"])
            if fill:
                ws.cell(ws.max_row, 2).fill = PatternFill("solid", fgColor=fill)
        for i, col in enumerate(cols, start=1):
            width = max([len(str(col))] + [min(len(str(r.get(col, ""))), 60) for r in rows])
            ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 62)
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions

        summary = wb.create_sheet("Сводка")
        summary.append(["Статус", "Количество"])
        summary["A1"].font = summary["B1"].font = Font(bold=True)
        counts: dict[str, int] = {}
        for r in results:
            counts[r.status] = counts.get(r.status, 0) + 1
        for status, n in sorted(counts.items(), key=lambda x: -x[1]):
            summary.append([status, n])
        summary.append(["Всего", len(results)])
        summary.column_dimensions["A"].width = 26

        wb.save(out_path)
        written.insert(0, out_path)
    return written


def load_previous(out_path: Path) -> dict[str, CheckResult]:
    """Для продолжения прерванной проверки: читаем уже проверенные УИН из CSV."""
    csv_path = out_path.with_suffix(".csv")
    done: dict[str, CheckResult] = {}
    if not csv_path.exists():
        return done
    with csv_path.open(encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f, delimiter=";"):
            if row.get("Статус") in (STATUS_ERROR, "", None):
                continue
            details = {
                k: v for k, v in row.items()
                if k not in BASE_COLUMNS and k not in TAIL_COLUMNS and v
            }
            done[row["УИН"]] = CheckResult(
                uin=row["УИН"],
                status=row["Статус"],
                status_raw=row.get("Статус (как на сайте)", ""),
                details=details,
                page_text=row.get("Текст ответа", ""),
                raw_json=row.get("JSON ответа", ""),
                error=row.get("Ошибка", ""),
                checked_at=row.get("Проверено", ""),
            )
    return done


# ---------------------------------------------------------------------------
# Основной цикл
# ---------------------------------------------------------------------------

def run(
    uins: list[str],
    out_path: Path,
    url: str = DEFAULT_URL,
    headless: bool = True,
    delay: float = 2.0,
    retries: int = 2,
    timeout: float = 30.0,
    resume: bool = True,
    input_selector: str | None = None,
    button_selector: str | None = None,
    executable_path: str | None = None,
    log: Callable[[str], None] = print,
    progress: Callable[[int, int], None] | None = None,
    should_stop: Callable[[], bool] | None = None,
) -> list[CheckResult]:
    previous = load_previous(out_path) if resume else {}
    results: dict[str, CheckResult] = {u: previous[u] for u in uins if u in previous}
    todo = [u for u in uins if u not in results]
    if results:
        log(f"Уже проверено ранее: {len(results)}, осталось: {len(todo)}")

    def ordered() -> list[CheckResult]:
        return [results[u] for u in uins if u in results]

    if todo:
        with DmdkChecker(
            url=url, headless=headless, timeout=timeout,
            input_selector=input_selector, button_selector=button_selector,
            executable_path=executable_path, log=log,
        ) as checker:
            for n, uin in enumerate(todo, start=1):
                if should_stop and should_stop():
                    log("Остановлено пользователем.")
                    break
                res: CheckResult | None = None
                for attempt in range(1, retries + 2):
                    try:
                        res = checker.check(uin)
                        if not res.error:
                            break
                    except Exception as e:  # noqa: BLE001 — любую ошибку сайта пишем в отчёт
                        res = CheckResult(uin=uin, status=STATUS_ERROR, error=str(e).splitlines()[0])
                    if attempt <= retries:
                        log(f"  повтор {attempt} для {uin}: {res.error}")
                        time.sleep(delay * 2)
                if res.error and res.status == STATUS_UNKNOWN:
                    res.status = STATUS_ERROR
                results[uin] = res
                log(f"[{len(results)}/{len(uins)}] {uin}: {res.status}"
                    + (f" ({res.status_raw})" if res.status_raw else "")
                    + (f" — {res.error}" if res.error else ""))
                if progress:
                    progress(len(results), len(uins))
                # Сохраняем после каждого УИН, чтобы ничего не потерять при сбое.
                save_results(ordered(), out_path)
                if n < len(todo):
                    time.sleep(delay)

    final = ordered()
    written = save_results(final, out_path)
    log("Готово. Файлы: " + ", ".join(str(p) for p in written))
    return final


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(
        description="Проверка списка УИН в ГИИС ДМДК (Пробирная палата РФ)",
    )
    p.add_argument("input", type=Path, help="Файл со списком УИН: .txt, .csv или .xlsx")
    p.add_argument("-o", "--output", type=Path,
                   help="Файл результата .xlsx (рядом будет .csv). По умолчанию results_<дата>.xlsx")
    p.add_argument("--url", default=DEFAULT_URL, help=f"Страница проверки (по умолчанию {DEFAULT_URL})")
    p.add_argument("--show-browser", action="store_true",
                   help="Показывать окно браузера (нужно, если сайт просит капчу)")
    p.add_argument("--delay", type=float, default=2.0, help="Пауза между запросами, сек (по умолчанию 2)")
    p.add_argument("--retries", type=int, default=2, help="Повторов при ошибке (по умолчанию 2)")
    p.add_argument("--timeout", type=float, default=30.0, help="Ожидание ответа сайта, сек")
    p.add_argument("--no-resume", action="store_true",
                   help="Проверить всё заново, не продолжая прошлый запуск")
    p.add_argument("--input-selector", help="CSS-селектор поля ввода УИН (если не находится сам)")
    p.add_argument("--button-selector", help="CSS-селектор кнопки «Проверить»")
    p.add_argument("--browser-path", help="Путь к Chrome/Chromium, если не установлен через playwright")
    args = p.parse_args(argv)

    if not args.input.exists():
        print(f"Файл не найден: {args.input}", file=sys.stderr)
        return 2
    uins = read_uins(args.input)
    if not uins:
        print("В файле не найдено ни одного УИН (ожидаются числовые коды).", file=sys.stderr)
        return 2
    out = args.output or Path(f"results_{datetime.now():%Y-%m-%d_%H-%M}.xlsx")
    print(f"УИН к проверке: {len(uins)}. Результат: {out}")
    run(
        uins, out, url=args.url, headless=not args.show_browser, delay=args.delay,
        retries=args.retries, timeout=args.timeout, resume=not args.no_resume,
        input_selector=args.input_selector, button_selector=args.button_selector,
        executable_path=args.browser_path,
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
